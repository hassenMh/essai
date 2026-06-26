from __future__ import annotations
from pathlib import Path
from datetime import datetime
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from reportlab.lib.pagesizes import A4
from reportlab.lib import colors
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
from reportlab.lib.styles import getSampleStyleSheet
from config import RECEIPTS_DIR


def export_to_excel(data: list[dict], filename: str, sheet_name: str = "Données") -> str:
    RECEIPTS_DIR.mkdir(parents=True, exist_ok=True)
    path = RECEIPTS_DIR / f"{filename}_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = sheet_name

    if not data:
        wb.save(str(path))
        return str(path)

    headers = list(data[0].keys())
    header_fill = PatternFill("solid", fgColor="2C3E50")
    header_font = Font(bold=True, color="FFFFFF")

    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=header.replace("_", " ").title())
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center")

    for row_idx, row_data in enumerate(data, 2):
        for col_idx, key in enumerate(headers, 1):
            ws.cell(row=row_idx, column=col_idx, value=row_data.get(key))

    for col in ws.columns:
        max_len = max(len(str(c.value or "")) for c in col)
        ws.column_dimensions[col[0].column_letter].width = min(max_len + 4, 40)

    wb.save(str(path))
    return str(path)


def generate_receipt_pdf(sale: dict, store_info: dict) -> str:
    RECEIPTS_DIR.mkdir(parents=True, exist_ok=True)
    filename = f"ticket_{sale['id']}_{datetime.now().strftime('%Y%m%d_%H%M%S')}.pdf"
    path = RECEIPTS_DIR / filename

    doc = SimpleDocTemplate(str(path), pagesize=A4, rightMargin=30, leftMargin=30, topMargin=20, bottomMargin=20)
    styles = getSampleStyleSheet()
    story = []

    # Header
    story.append(Paragraph(f"<b>{store_info.get('store_name', 'Magasin')}</b>", styles["Title"]))
    story.append(Paragraph(store_info.get("store_address", ""), styles["Normal"]))
    story.append(Paragraph(f"Tél: {store_info.get('store_phone', '')}", styles["Normal"]))
    story.append(Spacer(1, 12))
    story.append(Paragraph(f"Ticket N°: {sale['id']}", styles["Normal"]))
    story.append(Paragraph(f"Date: {sale['created_at']}", styles["Normal"]))
    story.append(Paragraph(f"Caissier: {sale.get('cashier_name', '')}", styles["Normal"]))
    story.append(Spacer(1, 12))

    # Items table
    table_data = [["Produit", "Qté", "P.U", "Total"]]
    for item in sale.get("items", []):
        table_data.append([
            item["product_name"],
            f"{item['quantity']:.2f}",
            f"{item['unit_price']:.3f}",
            f"{item['total']:.3f}",
        ])

    t = Table(table_data, colWidths=[200, 60, 80, 80])
    t.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#2C3E50")),
        ("TEXTCOLOR",  (0, 0), (-1, 0), colors.white),
        ("FONTNAME",   (0, 0), (-1, 0), "Helvetica-Bold"),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F8F9FA")]),
        ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
        ("ALIGN", (1, 0), (-1, -1), "RIGHT"),
    ]))
    story.append(t)
    story.append(Spacer(1, 12))

    # Totals
    currency = store_info.get("currency", "TND")
    story.append(Paragraph(f"Sous-total: {sale['subtotal']:.3f} {currency}", styles["Normal"]))
    if sale.get("discount", 0):
        story.append(Paragraph(f"Remise: -{sale['discount']:.3f} {currency}", styles["Normal"]))
    if sale.get("tax", 0):
        story.append(Paragraph(f"TVA: {sale['tax']:.3f} {currency}", styles["Normal"]))
    story.append(Paragraph(f"<b>TOTAL: {sale['total']:.3f} {currency}</b>", styles["Heading2"]))
    story.append(Paragraph(f"Payé: {sale['amount_paid']:.3f} {currency}", styles["Normal"]))
    story.append(Paragraph(f"Monnaie: {sale['change_given']:.3f} {currency}", styles["Normal"]))
    story.append(Spacer(1, 20))
    story.append(Paragraph(store_info.get("receipt_footer", "Merci !"), styles["Normal"]))

    doc.build(story)
    return str(path)
